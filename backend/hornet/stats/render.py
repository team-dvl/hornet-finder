"""
Files of a statistic, built from the table the API returns: CSV for any tool,
XLSX with a second sheet naming the period, filters and scope.

The CSV follows the Belgian and French spreadsheet convention (`;` between
fields, `,` in decimals, UTF-8 with a byte order mark) so it opens as a table
in Excel; the XLSX carries real numbers and dates.
"""

import csv
import io
from datetime import date

from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.utils.text import slugify

FORMATS = {
    'csv': 'text/csv; charset=utf-8',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}
GRANULARITY_LABELS = {'day': 'Par jour', 'week': 'Par semaine', 'month': 'Par mois'}
TOTAL_LABEL = 'Total'


def filename(result: dict, fmt: str) -> str:
    # The dash of a date range would otherwise vanish: "1–31 mars" is "1-31-mars"
    label = result['period']['label'].replace('–', ' ')
    stem = slugify(f"{result['statistic']['title']} {label}")
    return f"{stem}.{fmt}"


def _header(column: dict) -> str:
    return f"{column['label']} (%)" if column['type'] == 'percent' else column['label']


def _lines(result: dict) -> list:
    """The rows, then the totals under the first column's 'Total'."""
    totals = dict(result['totals'])
    first = result['columns'][0]['key']
    if result['columns'][0]['type'] == 'text':
        totals.setdefault(first, TOTAL_LABEL)
    if first == 'bucket':
        totals[first] = TOTAL_LABEL
        totals['start'], totals['end'] = result['period']['start'], result['period']['end']
    return result['rows'] + [totals]


def parameters(result: dict) -> list:
    period = result['period']
    computed = timezone.localtime(parse_datetime(result['computed_at']))
    lines = [
        ('Statistique', result['statistic']['title']),
        ('Période', f"{period['label']} ({period['dates']})"),
    ]
    if result.get('previous'):
        lines.append(('Comparée à', f"{result['previous']['label']} ({result['previous']['dates']})"))
    if result.get('granularity'):
        lines.append(('Granularité', GRANULARITY_LABELS[result['granularity']]))
    lines += [
        ('Filtres', ', '.join(result['filters']) or 'Aucun'),
        ('Pièges comptés', f"{result['scope']['label']} : {result['scope']['traps']}"),
        ('Calculé le', computed.strftime('%d/%m/%Y %H:%M')),
        ('Intervalles de confiance', 'À 95 % : loi de Poisson exacte pour les captures, '
                                     'Wilson pour la sélectivité'),
    ]
    lines += [tuple(note) for note in result.get('notes', [])]
    lines += [('Avertissement', warning) for warning in result['warnings']]
    return lines


def _text(value, column: dict) -> str:
    if value is None:
        return ''
    kind = column['type']
    if kind == 'int':
        return str(int(value))
    if kind == 'float':
        return f"{value:.{column.get('decimals', 2)}f}".replace('.', ',')
    if kind == 'percent':
        return f"{value * 100:.1f}".replace('.', ',')
    return str(value)


def render_csv(result: dict) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';', lineterminator='\r\n')
    columns = result['columns']
    writer.writerow([_header(c) for c in columns])
    for line in _lines(result):
        writer.writerow([_text(line.get(c['key']), c) for c in columns])
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def render_xlsx(result: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    formats = {'int': '0', 'percent': '0.0%', 'date': 'DD/MM/YYYY'}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Données'
    columns = result['columns']
    sheet.append([c['label'] for c in columns])
    for cell in sheet[1]:
        cell.font = bold
    lines = _lines(result)
    for index, line in enumerate(lines):
        values = []
        for column in columns:
            value = line.get(column['key'])
            if column['type'] == 'date' and value:
                value = date.fromisoformat(value)
            values.append(value)
        sheet.append(values)
        row = sheet.max_row
        for position, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row, column=position)
            number_format = formats.get(column['type'])
            if column['type'] == 'float':
                number_format = '0.' + '0' * column.get('decimals', 2)
            if number_format:
                cell.number_format = number_format
            if index == len(lines) - 1:
                cell.font = bold
    sheet.freeze_panes = 'A2'
    for position, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = max(12, len(column['label']) + 2)

    info = workbook.create_sheet('Paramètres')
    for label, value in parameters(result):
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = bold
    info.column_dimensions['A'].width = 26
    info.column_dimensions['B'].width = 90

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render(result: dict, fmt: str) -> bytes:
    return render_csv(result) if fmt == 'csv' else render_xlsx(result)
